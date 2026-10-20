"""Excel import of the price list and the doctor roster (ТЗ §11).

Column names are matched case-insensitively against several aliases, because
every clinic exports its price list differently.  Rows that cannot be parsed
are reported back rather than silently skipped — a missing service becomes an
invented price later, which is the one thing we never allow.
"""
from __future__ import annotations

import re
from typing import Any

from sqlalchemy import select
from sqlalchemy.ext.asyncio import AsyncSession

from arix.core.context import TenantContext
from arix.core.timeutil import parse_hhmm
from arix.verticals.clinic.models import Doctor, DoctorService, Schedule, Service

SERVICE_ALIASES = {
    "name": ("услуга", "наименование", "название", "service", "қызмет"),
    "category": ("категория", "раздел", "группа", "category"),
    "price_from": ("цена от", "цена", "от", "price", "price_from", "стоимость"),
    "price_to": ("цена до", "до", "price_to"),
    "duration_min": ("мин", "минут", "длительность", "duration", "время"),
    "buffer_min": ("буфер", "перерыв", "buffer"),
    "doctor": ("врач", "специалист", "doctor"),
    "recall_days": ("recall", "повтор", "контроль через"),
    "online": ("онлайн", "online", "онлайн-запись"),
}

DOCTOR_ALIASES = {
    "name": ("врач", "имя", "фио", "name", "doctor"),
    "specialty": ("специальность", "специализация", "specialty"),
    "langs": ("языки", "язык", "langs"),
    "services": ("услуги", "services"),
    "weekdays": ("дни", "график дни", "weekdays"),
    "start": ("начало", "с", "start"),
    "end": ("конец", "до", "end"),
    "break": ("обед", "перерыв", "break"),
}

WEEKDAY_WORDS = {
    "пн": 0, "вт": 1, "ср": 2, "чт": 3, "пт": 4, "сб": 5, "вс": 6,
    "mon": 0, "tue": 1, "wed": 2, "thu": 3, "fri": 4, "sat": 5, "sun": 6,
}


def _load_rows(path: str) -> list[dict[str, Any]]:
    from openpyxl import load_workbook

    workbook = load_workbook(path, data_only=True, read_only=True)
    sheet = workbook.active
    rows = list(sheet.iter_rows(values_only=True))
    workbook.close()
    if not rows:
        return []
    headers = [str(cell or "").strip().lower() for cell in rows[0]]
    out: list[dict[str, Any]] = []
    for row in rows[1:]:
        if all(cell is None or str(cell).strip() == "" for cell in row):
            continue
        out.append({headers[i]: row[i] for i in range(min(len(headers), len(row)))})
    return out


def _pick(row: dict[str, Any], aliases: tuple[str, ...]) -> Any:
    for key, value in row.items():
        if not key:
            continue
        for alias in aliases:
            if key == alias or key.startswith(alias) or alias in key:
                return value
    return None


def _number(value: Any) -> float | None:
    if value is None or value == "":
        return None
    if isinstance(value, (int, float)):
        return float(value)
    cleaned = re.sub(r"[^\d.,]", "", str(value)).replace(",", ".")
    try:
        return float(cleaned) if cleaned else None
    except ValueError:
        return None


def _slugify_code(name: str, used: set[str]) -> str:
    translit = str.maketrans(
        {
            "а": "a", "б": "b", "в": "v", "г": "g", "д": "d", "е": "e", "ж": "zh", "з": "z",
            "и": "i", "й": "y", "к": "k", "л": "l", "м": "m", "н": "n", "о": "o", "п": "p",
            "р": "r", "с": "s", "т": "t", "у": "u", "ф": "f", "х": "h", "ц": "c", "ч": "ch",
            "ш": "sh", "щ": "sch", "ы": "y", "э": "e", "ю": "yu", "я": "ya",
        }
    )
    base = re.sub(r"[^a-z0-9]+", "_", (name or "").lower().translate(translit)).strip("_")[:40]
    base = base or "service"
    code = base
    index = 2
    while code in used:
        code = f"{base}_{index}"
        index += 1
    used.add(code)
    return code


async def import_prices_xlsx(
    session: AsyncSession, ctx: TenantContext, path: str
) -> dict:
    rows = _load_rows(path)
    existing = {
        row.code: row
        for row in (
            await session.execute(select(Service).where(Service.tenant_id == ctx.id))
        ).scalars()
    }
    by_name = {
        (row.name or {}).get("ru", "").strip().lower(): row for row in existing.values()
    }
    used_codes = set(existing)
    created = updated = 0
    problems: list[str] = []

    for index, row in enumerate(rows, start=2):
        name = _pick(row, SERVICE_ALIASES["name"])
        if not name or not str(name).strip():
            problems.append(f"строка {index}: нет названия услуги")
            continue
        name = str(name).strip()
        price_from = _number(_pick(row, SERVICE_ALIASES["price_from"]))
        price_to = _number(_pick(row, SERVICE_ALIASES["price_to"]))
        duration = _number(_pick(row, SERVICE_ALIASES["duration_min"])) or 30
        buffer_min = _number(_pick(row, SERVICE_ALIASES["buffer_min"])) or 0
        recall = _number(_pick(row, SERVICE_ALIASES["recall_days"]))
        online_raw = _pick(row, SERVICE_ALIASES["online"])
        online = True
        if online_raw is not None and str(online_raw).strip() != "":
            online = str(online_raw).strip().lower() in ("1", "да", "yes", "true", "+", "y")

        service = by_name.get(name.lower())
        if service is None:
            service = Service(
                tenant_id=ctx.id,
                code=_slugify_code(name, used_codes),
                name={"ru": name},
                currency=ctx.currency,
            )
            session.add(service)
            created += 1
        else:
            updated += 1
        service.name = {**(service.name or {}), "ru": name}
        service.category = (
            str(_pick(row, SERVICE_ALIASES["category"]) or service.category or "").strip() or None
        )
        service.price_from = price_from
        service.price_to = price_to if price_to and price_to >= (price_from or 0) else None
        service.duration_min = int(duration)
        service.buffer_min = int(buffer_min)
        service.recall_days = int(recall) if recall else None
        service.online_bookable = online
        service.active = True
        if price_from is None:
            problems.append(f"строка {index}: услуга «{name}» без цены — бот не назовёт её цену")

    await session.flush()
    return {"created": created, "updated": updated, "rows": len(rows), "problems": problems}


async def import_doctors_xlsx(
    session: AsyncSession, ctx: TenantContext, path: str
) -> dict:
    rows = _load_rows(path)
    services = {
        (row.name or {}).get("ru", "").strip().lower(): row
        for row in (
            await session.execute(select(Service).where(Service.tenant_id == ctx.id))
        ).scalars()
    }
    existing = {
        row.name.strip().lower(): row
        for row in (
            await session.execute(select(Doctor).where(Doctor.tenant_id == ctx.id))
        ).scalars()
    }
    created = updated = schedules = 0
    problems: list[str] = []

    for index, row in enumerate(rows, start=2):
        name = _pick(row, DOCTOR_ALIASES["name"])
        if not name or not str(name).strip():
            problems.append(f"строка {index}: нет имени врача")
            continue
        name = str(name).strip()
        doctor = existing.get(name.lower())
        if doctor is None:
            doctor = Doctor(tenant_id=ctx.id, name=name)
            session.add(doctor)
            await session.flush()
            created += 1
        else:
            updated += 1
        doctor.specialty = (
            str(_pick(row, DOCTOR_ALIASES["specialty"]) or doctor.specialty or "").strip() or None
        )
        langs_raw = _pick(row, DOCTOR_ALIASES["langs"])
        if langs_raw:
            doctor.langs = [
                part.strip().lower()[:2] for part in re.split(r"[,;/]", str(langs_raw)) if part.strip()
            ]
        doctor.active = True

        services_raw = _pick(row, DOCTOR_ALIASES["services"])
        if services_raw:
            for part in re.split(r"[,;]", str(services_raw)):
                service = services.get(part.strip().lower())
                if service is None:
                    problems.append(
                        f"строка {index}: услуга «{part.strip()}» не найдена в прайсе"
                    )
                    continue
                already = (
                    await session.execute(
                        select(DoctorService).where(
                            DoctorService.doctor_id == doctor.id,
                            DoctorService.service_id == service.id,
                        )
                    )
                ).scalar_one_or_none()
                if already is None:
                    session.add(
                        DoctorService(
                            tenant_id=ctx.id, doctor_id=doctor.id, service_id=service.id
                        )
                    )

        weekdays_raw = _pick(row, DOCTOR_ALIASES["weekdays"])
        start_raw = _pick(row, DOCTOR_ALIASES["start"])
        end_raw = _pick(row, DOCTOR_ALIASES["end"])
        if weekdays_raw and start_raw and end_raw:
            weekdays = _parse_weekdays(str(weekdays_raw))
            break_raw = _pick(row, DOCTOR_ALIASES["break"])
            break_start = break_end = None
            if break_raw:
                parts = re.split(r"[-–—]", str(break_raw))
                if len(parts) == 2:
                    break_start = parse_hhmm(parts[0].strip())
                    break_end = parse_hhmm(parts[1].strip())
            for weekday in weekdays:
                exists = (
                    await session.execute(
                        select(Schedule).where(
                            Schedule.doctor_id == doctor.id, Schedule.weekday == weekday
                        )
                    )
                ).scalar_one_or_none()
                if exists is not None:
                    continue
                session.add(
                    Schedule(
                        tenant_id=ctx.id, doctor_id=doctor.id, weekday=weekday,
                        start=parse_hhmm(_time_text(start_raw)),
                        end=parse_hhmm(_time_text(end_raw)),
                        break_start=break_start, break_end=break_end,
                    )
                )
                schedules += 1

    await session.flush()
    return {
        "created": created, "updated": updated, "schedules": schedules,
        "rows": len(rows), "problems": problems,
    }


def _time_text(value: Any) -> str:
    if hasattr(value, "strftime"):
        return value.strftime("%H:%M")
    return str(value).strip()


def _parse_weekdays(raw: str) -> list[int]:
    raw = raw.lower().replace(" ", "")
    out: set[int] = set()
    for chunk in re.split(r"[,;]", raw):
        if not chunk:
            continue
        if "-" in chunk or "–" in chunk:
            parts = re.split(r"[-–]", chunk)
            if len(parts) == 2 and all(p[:2] in WEEKDAY_WORDS for p in parts):
                first, last = WEEKDAY_WORDS[parts[0][:2]], WEEKDAY_WORDS[parts[1][:2]]
                out.update(range(min(first, last), max(first, last) + 1))
                continue
        if chunk[:2] in WEEKDAY_WORDS:
            out.add(WEEKDAY_WORDS[chunk[:2]])
        elif chunk.isdigit() and 1 <= int(chunk) <= 7:
            out.add(int(chunk) - 1)
    return sorted(out)


__all__ = ["import_doctors_xlsx", "import_prices_xlsx"]
